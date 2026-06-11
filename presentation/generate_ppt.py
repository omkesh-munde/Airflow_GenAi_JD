"""
Generate an attractive, fully-editable PowerPoint deck for the
Airflow Log AI Analyzer project.

Run:
    python presentation/generate_ppt.py
Output:
    presentation/Airflow_Log_AI_Analyzer.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ---------------------------------------------------------------- theme
BG      = RGBColor(0x0F, 0x14, 0x19)
PANEL   = RGBColor(0x17, 0x21, 0x31)
PANEL2  = RGBColor(0x1F, 0x2C, 0x40)
BORDER  = RGBColor(0x2A, 0x3A, 0x52)
TEXT    = RGBColor(0xE7, 0xEC, 0xF3)
MUTED   = RGBColor(0x9A, 0xAA, 0xC0)
ACCENT  = RGBColor(0x3B, 0x82, 0xF6)
ACCENT2 = RGBColor(0x06, 0xB6, 0xD4)
GREEN   = RGBColor(0x22, 0xC5, 0x5E)
AMBER   = RGBColor(0xF5, 0x9E, 0x0B)
RED     = RGBColor(0xEF, 0x44, 0x44)
PURPLE  = RGBColor(0xA7, 0x8B, 0xFA)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Segoe UI"
FONT_MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def _no_line(shape):
    shape.line.fill.background()


def slide():
    s = prs.slides.add_slide(BLANK)
    rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    rect.fill.solid()
    rect.fill.fore_color.rgb = BG
    _no_line(rect)
    rect.shadow.inherit = False
    return s


def deco_corner(s):
    """Subtle accent shapes for visual interest."""
    c1 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-1.6), Inches(-1.8),
                            Inches(3.6), Inches(3.6))
    c1.fill.solid(); c1.fill.fore_color.rgb = RGBColor(0x14, 0x22, 0x3A)
    _no_line(c1); c1.shadow.inherit = False
    c2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.3), Inches(5.6),
                            Inches(3.4), Inches(3.4))
    c2.fill.solid(); c2.fill.fore_color.rgb = RGBColor(0x10, 0x26, 0x34)
    _no_line(c2); c2.shadow.inherit = False


def box(s, l, t, w, h, fill=None, line=None, rounded=False, radius=0.08,
        line_w=1.0):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        l, t, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        _no_line(shp)
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.shadow.inherit = False
    return shp


def text(s, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6, line_spacing=1.05):
    """runs: list of paragraphs; each paragraph is a list of (txt, size, color,
    bold, italic, font) tuples OR a single string."""
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        if isinstance(para, str):
            para = [(para, 18, TEXT, False, False, FONT)]
        for spec in para:
            txt, size, color, bold, italic, fnt = (list(spec) + [FONT])[:6]
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
            r.font.name = fnt; r.font.color.rgb = color
    return tb


def header(s, kicker, title, accent=ACCENT2):
    box(s, 0, 0, Inches(0.18), SH, fill=accent)              # left spine
    text(s, Inches(0.55), Inches(0.42), Inches(12), Inches(0.4),
         [[(kicker.upper(), 13, accent, True, False, FONT)]])
    text(s, Inches(0.52), Inches(0.78), Inches(12.3), Inches(0.9),
         [[(title, 32, TEXT, True, False, FONT)]])
    box(s, Inches(0.58), Inches(1.62), Inches(1.4), Inches(0.06), fill=accent)


def footer(s, n):
    text(s, Inches(0.55), Inches(7.04), Inches(9), Inches(0.35),
         [[("Airflow Log AI Analyzer  ·  Generic LLM + RAG", 9, MUTED,
            False, False, FONT)]])
    text(s, Inches(12.0), Inches(7.04), Inches(0.9), Inches(0.35),
         [[(f"{n:02d}", 10, MUTED, True, False, FONT)]], align=PP_ALIGN.RIGHT)


def bullet_list(s, l, t, w, h, items, size=16, gap=10, glyph="▸",
                glyph_color=ACCENT2, lead_color=TEXT):
    runs = []
    for it in items:
        if isinstance(it, tuple):
            head, rest = it
            runs.append([
                (f"{glyph}  ", size, glyph_color, True, False, FONT),
                (head, size, lead_color, True, False, FONT),
                (rest, size, MUTED, False, False, FONT),
            ])
        else:
            runs.append([
                (f"{glyph}  ", size, glyph_color, True, False, FONT),
                (it, size, TEXT, False, False, FONT),
            ])
    text(s, l, t, w, h, runs, space_after=gap, line_spacing=1.04)


def card(s, l, t, w, h, icon, title, body, accent):
    c = box(s, l, t, w, h, fill=PANEL, line=BORDER, rounded=True, radius=0.06)
    box(s, l, t, Inches(0.12), h, fill=accent)  # accent edge
    text(s, l + Inches(0.22), t + Inches(0.14), w - Inches(0.4), Inches(0.5),
         [[(f"{icon}  ", 18, accent, False, False, FONT),
           (title, 14.5, TEXT, True, False, FONT)]])
    text(s, l + Inches(0.24), t + Inches(0.62), w - Inches(0.42),
         h - Inches(0.72),
         [[(body, 11.5, MUTED, False, False, FONT)]], line_spacing=1.08)
    return c


# ================================================================ SLIDE 1
s = slide()
deco_corner(s)
box(s, 0, 0, SW, Inches(0.16), fill=ACCENT)
box(s, 0, Inches(0.16), SW, Inches(0.06), fill=ACCENT2)
text(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(0.5),
     [[("APACHE AIRFLOW  ·  GEN-AI AGENTS  ·  RAG MEMORY", 15, ACCENT2,
        True, False, FONT)]])
text(s, Inches(0.85), Inches(2.5), Inches(11.6), Inches(1.5),
     [[("Airflow Log AI Analyzer", 52, TEXT, True, False, FONT)]])
text(s, Inches(0.9), Inches(3.7), Inches(11.3), Inches(0.9),
     [[("Autonomous, RAG-powered root-cause analysis for Apache Airflow "
        "pipeline failures.", 20, MUTED, False, False, FONT)]])
# tech chips
chips = ["Python", "LangChain", "Chroma Vector DB", "SQLite", "Flask", "OpenAI / Ollama"]
cx = Inches(0.9)
for ch in chips:
    w = Inches(0.32 + 0.115 * len(ch))
    box(s, cx, Inches(4.85), w, Inches(0.46), fill=PANEL2, line=BORDER,
        rounded=True, radius=0.5)
    text(s, cx, Inches(4.85), w, Inches(0.46),
         [[(ch, 12, TEXT, False, False, FONT)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cx += w + Inches(0.18)
text(s, Inches(0.9), Inches(6.4), Inches(11), Inches(0.5),
     [[("Watches logs → classifies 15 error types → routes to specialist "
        "agents → remembers every fix.", 13, ACCENT2, False, True, FONT)]])

# ================================================================ SLIDE 2
s = slide(); deco_corner(s)
header(s, "The challenge", "Problem Statement", ACCENT2)
bullet_list(s, Inches(0.7), Inches(2.0), Inches(7.2), Inches(4.6), [
    ("Failures are diverse — ",
     "15+ distinct Airflow error modes (DAG imports, timeouts, scheduler, "
     "connections, XCom, OOM, Celery, K8s pods…)."),
    ("Logs are noisy — ",
     "long, scattered across scheduler, workers and tasks; engineers manually "
     "grep to find the real error."),
    ("Triage is slow & reactive — ",
     "diagnosis happens only after a pipeline breaks, often at 2 a.m."),
    ("Expertise bottleneck — ",
     "root-cause analysis depends on a few senior engineers."),
    ("No memory — ",
     "the same errors are re-investigated from scratch; lessons learned are lost."),
], size=15.5, gap=14)
# right stat panel
box(s, Inches(8.4), Inches(2.0), Inches(4.25), Inches(4.5), fill=PANEL,
    line=BORDER, rounded=True, radius=0.05)
text(s, Inches(8.7), Inches(2.25), Inches(3.7), Inches(0.5),
     [[("The cost of manual log triage", 14, TEXT, True, False, FONT)]])
stats = [("15+", "error categories to recognise", AMBER),
         ("High", "mean-time-to-resolution (MTTR)", RED),
         ("0", "institutional memory of past fixes", PURPLE),
         ("24/7", "uptime expected, but humans aren't", ACCENT2)]
yy = Inches(2.85)
for big, lab, col in stats:
    text(s, Inches(8.7), yy, Inches(3.7), Inches(0.55),
         [[(big + "   ", 26, col, True, False, FONT),
           (lab, 12.5, MUTED, False, False, FONT)]],
         anchor=MSO_ANCHOR.MIDDLE)
    yy += Inches(0.9)
footer(s, 2)

# ================================================================ SLIDE 3
s = slide(); deco_corner(s)
header(s, "Why today's approach falls short", "Current Limitations", AMBER)
lim = [
    ("Reactive, not proactive", "issues surface only after pipelines fail."),
    ("Human bottleneck", "depends on the availability & skill of on-call staff."),
    ("Inconsistent diagnoses", "quality varies by person, fatigue and time of day."),
    ("No reuse of past fixes", "zero shared memory across incidents."),
    ("Poor scalability", "can't keep up with more DAGs and growing log volume."),
    ("Alerts ≠ answers", "generic log tools notify, but don't explain root cause or fix."),
]
positions = [(0.7, 2.0), (4.75, 2.0), (8.8, 2.0),
             (0.7, 4.35), (4.75, 4.35), (8.8, 4.35)]
for (head, body), (lx, ty) in zip(lim, positions):
    card(s, Inches(lx), Inches(ty), Inches(3.75), Inches(2.05),
         "✕", head, body, AMBER)
footer(s, 3)

# ================================================================ SLIDE 4
s = slide(); deco_corner(s)
header(s, "An autonomous AI pipeline", "Proposed Solution", GREEN)
text(s, Inches(0.7), Inches(1.78), Inches(12), Inches(0.4),
     [[("extract → classify → route → retrieve → analyze → store → serve",
        14, ACCENT2, False, True, FONT_MONO)]])
sol = [
    ("👁", "Watch & detect", "watchdog auto-detects new log files in real time.", ACCENT),
    ("🏷", "Classify", "keyword engine maps each error to 15 Airflow categories.", ACCENT2),
    ("🧭", "Route to specialists", "Code / DB / Infra agents reason within their domain.", PURPLE),
    ("🧠", "RAG memory", "Chroma vector store recalls similar past incidents.", ACCENT2),
    ("⚙", "Context-aware LLM", "live pipeline config injected into every prompt.", GREEN),
    ("📋", "Structured output", "root cause + concrete fix + severity, every time.", AMBER),
]
pos = [(0.7, 2.35), (4.75, 2.35), (8.8, 2.35),
       (0.7, 4.55), (4.75, 4.55), (8.8, 4.55)]
for (ic, h, b, col), (lx, ty) in zip(sol, pos):
    card(s, Inches(lx), Inches(ty), Inches(3.75), Inches(1.95), ic, h, b, col)
footer(s, 4)

# ================================================================ SLIDE 5  (workflow)
s = slide(); deco_corner(s)
header(s, "How data flows end-to-end", "Workflow", ACCENT)
stages = [
    ("①", "Airflow DAG", "scheduled trigger", ACCENT),
    ("②", "Log Watcher", "watchdog detects log", ACCENT),
    ("③", "Extract", "find error block", ACCENT),
    ("④", "Classify", "15 error types", ACCENT2),
    ("⑤", "Route", "pick specialist", PURPLE),
    ("⑥", "Analyze", "LLM + RAG", GREEN),
    ("⑦", "Store", "SQLite + Chroma", AMBER),
    ("⑧", "Serve", "live dashboard", ACCENT2),
]
n = len(stages)
margin = Inches(0.6)
gap = Inches(0.12)
total_w = SW - 2 * margin
bw = (total_w - gap * (n - 1)) // n
bh = Inches(1.5)
ty = Inches(2.25)
centers = []
for i, (num, h, b, col) in enumerate(stages):
    lx = margin + i * (bw + gap)
    box(s, lx, ty, bw, bh, fill=PANEL, line=col, rounded=True, radius=0.08,
        line_w=1.5)
    box(s, lx, ty, bw, Inches(0.1), fill=col)
    text(s, lx, ty + Inches(0.16), bw, Inches(0.5),
         [[(num, 22, col, True, False, FONT)]], align=PP_ALIGN.CENTER)
    text(s, lx + Inches(0.04), ty + Inches(0.66), bw - Inches(0.08), Inches(0.4),
         [[(h, 12.5, TEXT, True, False, FONT)]], align=PP_ALIGN.CENTER)
    text(s, lx + Inches(0.04), ty + Inches(1.02), bw - Inches(0.08), Inches(0.4),
         [[(b, 9.5, MUTED, False, False, FONT)]], align=PP_ALIGN.CENTER)
    centers.append((lx, lx + bw))
    if i < n - 1:
        ax = lx + bw + Inches(0.005)
        arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax,
                                 ty + bh / 2 - Inches(0.08),
                                 gap - Inches(0.01), Inches(0.16))
        arr.fill.solid(); arr.fill.fore_color.rgb = ACCENT2
        _no_line(arr); arr.shadow.inherit = False

# supporting engines feeding "Analyze"
eng_y = Inches(4.55)
engines = [
    ("🧠  RAG Memory", "memory/store.py · retrieve_similar()", ACCENT2),
    ("🗄  Vector DB", "Chroma — embeddings of past errors", ACCENT2),
    ("🤖  LLM Engine", "LangChain → OpenAI / Ollama / Anthropic", GREEN),
    ("⚙  Pipeline Context", "config.py + *.toml injected to prompt", PURPLE),
]
ew = Inches(2.95); eg = Inches(0.15); ex = margin
text(s, Inches(0.6), Inches(4.1), Inches(12), Inches(0.4),
     [[("Engines feeding the Analyze stage", 13, MUTED, True, False, FONT)]])
for h, b, col in engines:
    box(s, ex, eng_y, ew, Inches(1.15), fill=PANEL2, line=BORDER, rounded=True,
        radius=0.08)
    text(s, ex + Inches(0.18), eng_y + Inches(0.16), ew - Inches(0.3), Inches(0.4),
         [[(h, 13.5, col, True, False, FONT)]])
    text(s, ex + Inches(0.18), eng_y + Inches(0.6), ew - Inches(0.3), Inches(0.45),
         [[(b, 10.5, MUTED, False, False, FONT_MONO)]])
    ex += ew + eg
text(s, Inches(0.6), Inches(6.0), Inches(12), Inches(0.4),
     [[("Every analysis is written back to the Vector DB — the system gets "
        "smarter with each incident.", 12, ACCENT2, False, True, FONT)]])
footer(s, 5)

# ================================================================ SLIDE 6
s = slide(); deco_corner(s)
header(s, "What you gain", "Benefits & Value Addition", GREEN)
left = [
    ("Faster MTTR — ", "instant root cause + fix instead of manual log hunting."),
    ("Always-on — ", "24/7 autonomous monitoring with no human in the loop."),
    ("Consistent & explainable — ", "structured root cause, fix & severity every time."),
    ("Institutional memory — ", "every incident improves future answers via RAG."),
    ("Full audit trail — ", "searchable history, live dashboard & value report."),
]
right = [
    ("Provider-agnostic — ", "OpenAI, Anthropic, Google, Ollama, Groq, Azure, HF."),
    ("Context-aware — ", "tuned to your Airflow version, executor & environment."),
    ("Scales effortlessly — ", "handles many DAGs and high log volume."),
    ("Low / zero cost — ", "runs locally; optional local LLM via Ollama."),
    ("Easy to extend — ", "config-driven error types, agents & prompts (TOML)."),
]
bullet_list(s, Inches(0.7), Inches(2.05), Inches(6.0), Inches(4.6), left,
            size=15, gap=15, glyph="✓", glyph_color=GREEN)
bullet_list(s, Inches(6.9), Inches(2.05), Inches(6.0), Inches(4.6), right,
            size=15, gap=15, glyph="✓", glyph_color=GREEN)
footer(s, 6)

# ================================================================ SLIDE 7  (table)
s = slide(); deco_corner(s)
header(s, "Manual monitoring vs. this AI agent", "Comparison", ACCENT2)
rows = [
    ("Dimension", "Manual Monitoring", "AI Agent (Automated)"),
    ("Detection", "Reactive — after failure is noticed", "Real-time on log creation"),
    ("Triage speed", "Minutes to hours per incident", "Seconds, fully automatic"),
    ("Root-cause analysis", "Manual log reading, senior-dependent", "LLM + RAG, structured & instant"),
    ("Consistency", "Varies by person & fatigue", "Uniform, repeatable output"),
    ("Memory of past fixes", "Tribal knowledge, often lost", "Persistent vector + SQLite memory"),
    ("Coverage", "Business hours / on-call", "24/7 unattended"),
    ("Expertise required", "High — senior engineers", "Low — expertise encoded in agents"),
    ("Scalability", "Poor — human bottleneck", "High — scales with compute"),
    ("Cost over time", "Grows with team & incidents", "Flat & low; local LLM ≈ zero"),
    ("Output", "Ad-hoc notes", "Root cause + fix + severity + audit"),
]
nrows, ncols = len(rows), 3
tbl_l, tbl_t = Inches(0.6), Inches(1.95)
tbl_w, tbl_h = Inches(12.1), Inches(4.85)
gf = s.shapes.add_table(nrows, ncols, tbl_l, tbl_t, tbl_w, tbl_h)
table = gf.table
table.columns[0].width = Inches(2.9)
table.columns[1].width = Inches(4.4)
table.columns[2].width = Inches(4.8)
# disable banding style by setting our own fills
table.first_row = False
table.horz_banding = False
for r in range(nrows):
    table.rows[r].height = Inches(0.41)
    for c in range(ncols):
        cell = table.cell(r, c)
        cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        if r == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = ACCENT
        elif c == 2:
            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0x12, 0x2A, 0x2E)
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if r % 2 else PANEL2
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        run = p.add_run(); run.text = rows[r][c]
        run.font.name = FONT
        if r == 0:
            run.font.size = Pt(13); run.font.bold = True; run.font.color.rgb = WHITE
        else:
            run.font.size = Pt(11.5)
            if c == 0:
                run.font.bold = True; run.font.color.rgb = TEXT
            elif c == 1:
                run.font.color.rgb = MUTED
            else:
                run.font.bold = True; run.font.color.rgb = GREEN
footer(s, 7)

# ================================================================ SLIDE 8
s = slide(); deco_corner(s)
box(s, 0, Inches(3.5), SW, Inches(0.06), fill=ACCENT2)
text(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.0),
     [[("From reactive firefighting to autonomous insight.", 30, TEXT, True,
        False, FONT)]])
text(s, Inches(0.92), Inches(3.7), Inches(11.5), Inches(0.6),
     [[("Run it in minutes:", 15, ACCENT2, True, False, FONT)]])
text(s, Inches(0.95), Inches(4.2), Inches(11.5), Inches(1.4),
     [[("pip install -r requirements.txt", 14, TEXT, False, False, FONT_MONO)],
      [("python main.py --simulate --simulate-delay 0.2", 14, TEXT, False,
        False, FONT_MONO)],
      [("→  open http://127.0.0.1:8080  (Dashboard + View Workflow)", 14,
        GREEN, False, False, FONT_MONO)]],
     space_after=8)
text(s, Inches(0.9), Inches(6.2), Inches(11.5), Inches(0.5),
     [[("Python · LangChain · Chroma · SQLite · Flask · OpenAI / Ollama",
        12, MUTED, False, False, FONT)]])

# ---------------------------------------------------------------- save
out = Path(__file__).resolve().parent / "Airflow_Log_AI_Analyzer.pptx"
prs.save(out)
print(f"Saved presentation: {out}")
print(f"Slides: {len(prs.slides._sldIdLst)}")
